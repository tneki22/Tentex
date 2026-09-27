"""DOCX: формулы Word и таблицы не пропадают.

`python-docx` не видит `<m:oMath>` в тексте абзаца и таблиц в
`document.paragraphs`: методичка из Word приезжала без формул и без таблиц.
"""

from docx import Document
from docx.oxml import parse_xml

from app.materials.parsers import native
from app.materials.parsers.docx_content import omml_latex, table_markdown

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _math(body: str, tag: str = "oMath") -> str:
    return f'<m:{tag} xmlns:m="{M_NS}" xmlns:w="{W_NS}">{body}</m:{tag}>'


def _r(text: str) -> str:
    return f"<m:r><m:t>{text}</m:t></m:r>"


def test_omml_structures_become_latex() -> None:
    fraction = parse_xml(_math(f"<m:f><m:num>{_r('1')}</m:num><m:den>{_r('n')}</m:den></m:f>"))
    assert omml_latex(fraction) == r"\frac{1}{n}"

    power = parse_xml(_math(f"<m:sSup><m:e>{_r('x')}</m:e><m:sup>{_r('2')}</m:sup></m:sSup>"))
    assert omml_latex(power) == "{x}^{2}"

    negation = parse_xml(
        _math(
            f'<m:bar><m:barPr><m:pos m:val="top"/></m:barPr><m:e>{_r("A")}</m:e></m:bar>'
            f"{_r('∧B≡C')}"
        )
    )
    assert omml_latex(negation) == r"\overline{A}\wedge B\equiv C"

    total = parse_xml(
        _math(
            '<m:nary><m:naryPr><m:chr m:val="∑"/></m:naryPr>'
            f"<m:sub>{_r('i=1')}</m:sub><m:sup>{_r('n')}</m:sup><m:e>{_r('x')}</m:e></m:nary>"
        )
    )
    assert omml_latex(total) == r"\sum_{i=1}^{n} x"

    root = parse_xml(_math(f"<m:rad><m:deg/><m:e>{_r('a')}</m:e></m:rad>"))
    assert omml_latex(root) == r"\sqrt{a}"

    braces = parse_xml(
        _math(
            '<m:d><m:dPr><m:begChr m:val="{"/><m:endChr m:val=""/></m:dPr>'
            f"<m:e>{_r('x')}</m:e></m:d>"
        )
    )
    assert omml_latex(braces) == r"\left\{x\right."

    words = parse_xml(_math(_r("если") + _r("x&gt;0")))
    assert omml_latex(words) == r"\text{если}x>0"


def test_table_keeps_formula_cells_and_merged_columns() -> None:
    cell = (
        '<w:tc><w:p><m:oMath xmlns:m="' + M_NS + '"><m:r><m:t>|x|</m:t></m:r></m:oMath>'
        "</w:p></w:tc>"
    )
    wide = (
        '<w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr>'
        "<w:p><w:r><w:t>Итог | всего</w:t></w:r></w:p></w:tc>"
    )
    table = parse_xml(
        f'<w:tbl xmlns:w="{W_NS}">'
        "<w:tr><w:tc><w:p><w:r><w:t>A</w:t></w:r></w:p></w:tc>"
        "<w:tc><w:p><w:r><w:t>B</w:t></w:r></w:p></w:tc></w:tr>"
        f"<w:tr>{cell}<w:tc><w:p/></w:tc></w:tr>"
        f"<w:tr>{wide}</w:tr>"
        "</w:tbl>"
    )

    assert table_markdown(table) == "\n".join(
        [
            "| A | B |",
            "|---|---|",
            "| $|x|$ |  |",
            r"| Итог \| всего |  |",
        ]
    )


def test_docx_page_reads_formulas_tables_in_document_order(tmp_path) -> None:
    document = Document()
    inline = document.add_paragraph("Пусть ")
    inline._p.append(parse_xml(_math(_r("A,B,C"))))
    inline.add_run(" – произвольные формулы.")
    display = document.add_paragraph()
    display._p.append(parse_xml(_math(_math(_r("F≡F")), "oMathPara")))
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "x"
    table.cell(0, 1).text = "¬x"
    table.cell(1, 0).text = "0"
    table.cell(1, 1).text = "1"
    document.add_paragraph("После таблицы.")
    path = tmp_path / "math.docx"
    document.save(path)

    parsed = native._docx_page(path)

    assert [element.kind for element in parsed.elements] == [
        "paragraph", "formula", "table", "paragraph",
    ]
    assert parsed.elements[0].text == "Пусть $A,B,C$ – произвольные формулы."
    assert parsed.elements[1].text == r"$$F\equiv F$$"
    assert parsed.elements[2].text.startswith("| x | ¬x |\n|---|---|\n| 0 | 1 |")
    assert parsed.elements[3].text == "После таблицы."
