import logging
import re
from collections.abc import Iterator, Sequence
from dataclasses import dataclass, replace
from difflib import SequenceMatcher
from functools import reduce
from pathlib import Path
from statistics import median
from tempfile import NamedTemporaryFile

import pymupdf as fitz
from docx import Document
from docx.text.paragraph import Paragraph
from PIL import Image

from app.materials.parsers import paddle_fast, raster, reading_order
from app.materials.parsers.audio import parse_audio
from app.materials.parsers.base import (
    IMAGE_PLACEHOLDER,
    ElementKind,
    PageRecognizer,
    ParsedElement,
    ParsedPage,
    RecognitionSource,
    RegionRequest,
)
from app.materials.parsers.cloud_vlm import wrap_bare_latex
from app.materials.parsers.pdf_layout import parse_layout_page
from app.materials.storage import store_material_asset
from app.models import ParserMode
from app.ocr.engines import DEFAULT_QUALITY_THRESHOLD, OcrRuntimeParams

NUMBERED_RE = re.compile(r"^\s*\d{1,3}[.)]\s+")
BULLET_RE = re.compile(r"^\s*[•▪◦‣·*+\-–—]\s+")
ORDERED_RE = re.compile(r"^\s*(?:\d{1,3}|[a-zа-я])[.)]\s+")
# Заголовок из одной строки: длиннее — это уже пункт списка, а не название вопроса.
HEADING_MAX_CHARS = 120
# Шаг отступа одного уровня вложенности, в пунктах PDF.
INDENT_STEP = 18
MAX_LIST_LEVEL = 4
# Насколько строка может не дотянуть до правого края и всё ещё считаться перенесённой.
WRAP_TOLERANCE = 60
ENDS_SENTENCE_RE = re.compile(r"[.!?:;][\"»)\]]?$")
# Кегль заголовка относительно основного текста.
HEADING_SIZE_RATIO = 1.15
# Меньше — это логотипы, линейки и артефакты вёрстки, а не иллюстрации.
MIN_IMAGE_SIDE = 40
# Виды элементов, у которых оригинальная вырезка со страницы полезна сама по
# себе: распознанному тексту формулы или схемы верить нельзя, и просмотрщик
# показывает рядом исходник.
CROPPED_KINDS = {"image", "formula", "table"}

log = logging.getLogger("tentex.worker")


def extract_outline(path: Path) -> list[dict[str, object]]:
    if path.suffix.lower() != ".pdf":
        return []
    with fitz.open(path) as document:
        return [
            {"level": level, "title": title.strip(), "page": page}
            for level, title, page, *_ in document.get_toc(simple=False)
            if title.strip() and page > 0
        ]


def inspect(path: Path) -> tuple[int, int, int, list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            document = fitz.open(path)
        except Exception as error:
            raise ValueError("PDF повреждён или не читается") from error
        if document.needs_pass:
            raise PermissionError("PDF защищён паролем")
        if len(document) > 500:
            raise OverflowError("В PDF больше 500 страниц")
        scan_pages = sum(not page.get_text("text").strip() for page in document)
        return len(document), scan_pages, max(1, len(document) - scan_pages + scan_pages * 3), []
    if suffix in {".jpg", ".jpeg", ".png"}:
        with Image.open(path) as image:
            image.verify()
        return 1, 1, 3, []
    if suffix == ".docx":
        Document(path)
        return 1, 0, 1, []
    if suffix in {".txt", ".md"}:
        path.read_text(encoding="utf-8-sig")
        return 1, 0, 1, []
    if suffix in {".mp3", ".wav", ".m4a", ".ogg", ".flac"}:
        return 1, 0, 60, ["audio_transcription_required"]
    raise ValueError("Неподдерживаемый формат")


def _normalized_bbox(
    bbox: tuple[float, float, float, float], width: float, height: float
) -> tuple[float, float, float, float]:
    x0, y0, x1, y1 = bbox
    return (max(0, x0 / width), max(0, y0 / height), min(1, x1 / width), min(1, y1 / height))


@dataclass(frozen=True, slots=True)
class _Line:
    text: str
    x0: float
    x1: float
    top: float
    bottom: float
    size: float


def _block_lines(block: dict) -> list[_Line]:
    lines: list[_Line] = []
    for line in block.get("lines", []):
        spans = [span for span in line.get("spans", []) if span.get("text", "").strip()]
        if not spans:
            continue
        text = "".join(str(span["text"]) for span in spans).strip()
        if not text:
            continue
        x0, top, x1, bottom = line["bbox"]
        lines.append(
            _Line(text, x0, x1, top, bottom, max(float(s.get("size", 0)) for s in spans))
        )
    return lines


def _join_wrapped(previous: str, addition: str) -> str:
    """Строки PDF — это перенос, а не конец абзаца. Слово, разорванное дефисом, склеивается."""
    if previous.endswith("-") and addition[:1].islower():
        return previous[:-1] + addition
    return f"{previous} {addition}"


def _continues(previous: _Line, line: _Line, wrap_threshold: float, heading_size: float) -> bool:
    """Продолжает ли строка предыдущую.

    Одной геометрии мало: текст не выключен по формату, и правый край гуляет на
    полсотни пунктов. Одной пунктуации тоже мало: «…(INSERT, UPDATE, DELETE — при»
    обязана склеиться со следующей строкой. Поэтому три признака вместе.
    """
    if BULLET_RE.match(line.text) or ORDERED_RE.match(line.text):
        return False
    if previous.size >= heading_size and line.size >= heading_size:
        return True  # многострочный заголовок: «Вопрос 21» и его название
    if ENDS_SENTENCE_RE.search(previous.text):
        return False
    return previous.x1 >= wrap_threshold


def _group_lines(lines: list[_Line], heading_size: float) -> list[list[_Line]]:
    if not lines:
        return []
    wrap_threshold = max(line.x1 for line in lines) - WRAP_TOLERANCE
    groups: list[list[_Line]] = [[lines[0]]]
    for previous, line in zip(lines, lines[1:], strict=False):
        if _continues(previous, line, wrap_threshold, heading_size):
            groups[-1].append(line)
        else:
            groups.append([line])
    return groups


def _indent_level(x0: float, base_x0: float) -> int:
    return 1 + min(MAX_LIST_LEVEL - 1, max(0, int((x0 - base_x0) / INDENT_STEP)))


def _classify(
    text: str, group: list[_Line], body_size: float, base_x0: float
) -> tuple[ElementKind, int | None]:
    max_size = max(line.size for line in group)
    bullet = BULLET_RE.match(text) is not None
    if not bullet:
        big = max_size >= body_size * HEADING_SIZE_RATIO
        short_numbered = (
            NUMBERED_RE.match(text) is not None
            and len(group) == 1
            and len(text) <= HEADING_MAX_CHARS
        )
        if big or short_numbered:
            return "heading", max(1, round(body_size * 2 / max_size))
    if bullet or ORDERED_RE.match(text) is not None:
        return "list", _indent_level(min(line.x0 for line in group), base_x0)
    return "paragraph", None


def _image_element(
    block: dict,
    page: fitz.Page,
    page_number: int,
    owner: str,
    index: int,
    *,
    run_ocr: bool = False,
    params: OcrRuntimeParams | None = None,
) -> ParsedElement | None:
    """Фотография или схема со страницы. Мелочь вроде логотипов пропускаем."""
    data = block.get("image")
    if not data:
        return None
    params = params or OcrRuntimeParams()
    x0, top, x1, bottom = block["bbox"]
    if (x1 - x0) < MIN_IMAGE_SIDE or (bottom - top) < MIN_IMAGE_SIDE:
        return None
    extension = str(block.get("ext") or "png").lower()
    asset_path = store_material_asset(owner, f"p{page_number}-{index}.{extension}", data)
    text = IMAGE_PLACEHOLDER
    confidence = None
    recognition_source: RecognitionSource = "native"
    if run_ocr and paddle_fast.available():
        with NamedTemporaryFile(suffix=f".{extension}", delete=False) as temporary:
            temporary.write(data)
            temporary_path = Path(temporary.name)
        try:
            recognized = paddle_fast.parse_image(
                temporary_path,
                page_number,
                language=params.fast_language,
                ocr_version=params.fast_model_id,
                quality_threshold=params.quality_threshold,
            )
            if recognized.plain_text.strip():
                text = recognized.plain_text.strip()
                confidence = recognized.confidence
                recognition_source = "ocr"
        except (OSError, RuntimeError, ValueError) as error:
            # Исходный вырез остаётся полезным, даже если OCR этой области не
            # справился, — но молчать об этом нельзя: иначе пропажа текста
            # картинки выглядит как «так и было».
            log.warning(
                "OCR картинки не удался page=%s index=%s: %s", page_number, index, error
            )
        finally:
            temporary_path.unlink(missing_ok=True)
    return ParsedElement(
        "image",
        text,
        _normalized_bbox((x0, top, x1, bottom), page.rect.width, page.rect.height),
        None,
        confidence,
        asset_path=asset_path,
        recognition_source=recognition_source,
    )


def _native_pdf_page(
    page: fitz.Page,
    page_number: int,
    owner: str = "",
    *,
    ocr_images: bool = False,
    params: OcrRuntimeParams | None = None,
) -> ParsedPage:
    raw = page.get_text("dict", sort=True)
    blocks = raw.get("blocks", [])
    text_blocks = [block for block in blocks if block.get("type") == 0]
    all_lines = [line for block in text_blocks for line in _block_lines(block)]
    if not all_lines:
        return ParsedPage(page_number, page.rect.width, page.rect.height, "", "", "native", ())
    body_size = median([line.size for line in all_lines])
    base_x0 = min(line.x0 for line in all_lines)
    heading_size = body_size * HEADING_SIZE_RATIO

    elements: list[ParsedElement] = []
    for index, block in enumerate(blocks):
        if block.get("type") == 1:
            if owner:
                image = _image_element(
                    block,
                    page,
                    page_number,
                    owner,
                    index,
                    run_ocr=ocr_images,
                    params=params,
                )
                if image is not None:
                    elements.append(image)
            continue
        for group in _group_lines(_block_lines(block), heading_size):
            text = reduce(_join_wrapped, (line.text for line in group))
            if not text.strip():
                continue
            kind, level = _classify(text, group, body_size, base_x0)
            bbox = (
                min(line.x0 for line in group),
                min(line.top for line in group),
                max(line.x1 for line in group),
                max(line.bottom for line in group),
            )
            elements.append(
                ParsedElement(
                    kind,
                    text,
                    _normalized_bbox(bbox, page.rect.width, page.rect.height),
                    level,
                )
            )

    # Картинка в plain-тексте была бы шумом: она попадает и в поиск, и в эталон.
    plain = "\n".join(element.text for element in elements if element.kind != "image")
    markdown = _markdown(elements)
    diagnostics = ("formula_possible",) if re.search(r"[∑∫√≈≤≥]", plain) else ()
    return ParsedPage(
        page_number,
        page.rect.width,
        page.rect.height,
        markdown,
        plain,
        "native",
        tuple(elements),
        diagnostics,
    )


def text_layer_pages(path: Path, owner: str = "") -> Iterator[ParsedPage]:
    """Страницы PDF, у которого текстовый слой заведомо полный и точный.

    Так читается PDF, который Tentex собрал сам (Typst): распознавать в нём
    нечего, а разметчик стоит времени и модели. Разбор всё равно идёт через
    общий `_native_pdf_page`, поэтому заголовки, списки и вырезы иллюстраций
    получаются те же, что у обычного PDF, — а значит работают и блоки, и
    привязки, и автопривязка ответов.
    """
    with fitz.open(path) as document:
        for index, page in enumerate(document):
            yield _native_pdf_page(page, index + 1, owner, ocr_images=False)


def _bbox_overlap(
    inner: tuple[float, float, float, float], outer: tuple[float, float, float, float]
) -> float:
    x0 = max(inner[0], outer[0])
    y0 = max(inner[1], outer[1])
    x1 = min(inner[2], outer[2])
    y1 = min(inner[3], outer[3])
    intersection = max(0.0, x1 - x0) * max(0.0, y1 - y0)
    area = max(0.000001, (inner[2] - inner[0]) * (inner[3] - inner[1]))
    return intersection / area


def _normalized_text(text: str) -> str:
    return re.sub(r"[^0-9a-zа-я]+", "", text.casefold())


def _merge_native_and_images(
    native: tuple[ParsedElement, ...], images: tuple[ParsedElement, ...]
) -> tuple[ParsedElement, ...]:
    """Merge layout regions and raster regions, preferring a native text layer."""
    merged_images: list[ParsedElement] = []
    native_text = [item for item in native if item.kind != "image" and item.text.strip()]
    for image in images:
        if image.recognition_source != "ocr":
            merged_images.append(image)
            continue
        covered = " ".join(
            item.text for item in native_text if _bbox_overlap(item.bbox, image.bbox) >= 0.7
        )
        left = _normalized_text(covered)
        right = _normalized_text(image.text)
        duplicate = bool(left and right) and (
            left in right
            or right in left
            or SequenceMatcher(None, left, right).ratio() >= 0.72
        )
        merged_images.append(
            replace(
                image,
                text=IMAGE_PLACEHOLDER,
                confidence=None,
                recognition_source="native",
            )
            if duplicate
            else image
        )
    return _in_reading_order((*native, *merged_images))


def _in_reading_order(elements: tuple[ParsedElement, ...]) -> tuple[ParsedElement, ...]:
    """Расставить элементы страницы в порядке чтения.

    Сортировка «сверху вниз, слева направо» здесь была бы ошибкой: на статье в
    две колонки она перемешивает колонки, ради чего и заведён `reading_order`.
    """
    order = reading_order.reading_order([element.bbox for element in elements])
    return tuple(elements[index] for index in order)


def _page_quality(
    elements: tuple[ParsedElement, ...],
    quality_threshold: float = DEFAULT_QUALITY_THRESHOLD,
) -> tuple[str, float | None]:
    recognized = [
        item
        for item in elements
        if item.recognition_source in {"ocr", "vl"} and item.text.strip()
    ]
    if not recognized:
        return "native", None
    scores = [item.confidence for item in recognized if item.confidence is not None]
    confidence = min(scores) if scores else None
    quality = "ocr_low" if confidence is not None and confidence < quality_threshold else "ocr"
    return quality, confidence


def _markdown(elements: list[ParsedElement]) -> str:
    """Разметка сохраняет уровни: заголовки решётками, пункты списка — отступом."""
    lines: list[str] = []
    for element in elements:
        if element.kind == "heading":
            lines.append(f"{'#' * (element.level or 2)} {element.text}")
        elif element.kind == "list":
            lines.append(f"{'    ' * ((element.level or 1) - 1)}{element.text}")
        elif element.kind == "image":
            lines.append(f"![Изображение]({element.asset_path or ''})")
        else:
            lines.append(element.text)
    return "\n\n".join(lines)


def parse_text_page(
    text: str,
    page_number: int = 1,
    recognition_source: RecognitionSource = "native",
) -> ParsedPage:
    elements: list[ParsedElement] = []
    raw_lines = [line for line in text.splitlines() if line.strip()]
    for index, raw in enumerate(raw_lines):
        line = raw.strip()
        markdown_heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        kind: ElementKind
        if markdown_heading:
            kind, content, level = (
                "heading",
                markdown_heading.group(2),
                len(markdown_heading.group(1)),
            )
        elif BULLET_RE.match(line):
            # Отступ в исходнике — это и есть уровень вложенности.
            indent = len(raw) - len(raw.lstrip())
            kind, content, level = "list", line, 1 + min(MAX_LIST_LEVEL - 1, indent // 2)
        elif NUMBERED_RE.match(line) and len(line) <= HEADING_MAX_CHARS:
            kind, content, level = "heading", line, 1
        elif ORDERED_RE.match(line):
            indent = len(raw) - len(raw.lstrip())
            kind, content, level = "list", line, 1 + min(MAX_LIST_LEVEL - 1, indent // 2)
        else:
            kind, content, level = "paragraph", line, None
        top = index / max(1, len(raw_lines))
        bottom = (index + 1) / max(1, len(raw_lines))
        elements.append(
            ParsedElement(
                kind,
                content,
                (0, top, 1, bottom),
                level,
                recognition_source=recognition_source,
            )
        )
    plain = "\n".join(element.text for element in elements)
    return ParsedPage(page_number, 1, 1, _markdown(elements), plain, "native", tuple(elements))


def _text_page(path: Path) -> ParsedPage:
    return parse_text_page(path.read_text(encoding="utf-8-sig"))


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
BULLET_FORMATS = {"bullet", "none"}
LETTER_FORMATS = {
    "lowerLetter": "абвгдежзиклмнопрстуфхцчшщэюя",
    "upperLetter": "АБВГДЕЖЗИКЛМНОПРСТУФХЦЧШЩЭЮЯ",
}


def _docx_numbering(document: Document) -> dict[tuple[int, int], tuple[str, str]]:
    """(numId, ilvl) → (формат, шаблон) из numbering.xml.

    Word держит номера пунктов не в тексте абзаца, а в отдельной части документа.
    Без этой карты нумерованный список приезжает без номеров — а в списке
    вопросов номер и есть половина смысла.
    """
    try:
        element = document.part.numbering_part.element
    except (KeyError, AttributeError, NotImplementedError):
        return {}

    abstract: dict[str, dict[int, tuple[str, str]]] = {}
    for node in element.findall(f"{W_NS}abstractNum"):
        levels: dict[int, tuple[str, str]] = {}
        for level in node.findall(f"{W_NS}lvl"):
            fmt = level.find(f"{W_NS}numFmt")
            template = level.find(f"{W_NS}lvlText")
            levels[int(level.get(f"{W_NS}ilvl", "0"))] = (
                fmt.get(f"{W_NS}val", "decimal") if fmt is not None else "decimal",
                template.get(f"{W_NS}val", "%1.") if template is not None else "%1.",
            )
        abstract[node.get(f"{W_NS}abstractNumId", "")] = levels

    mapping: dict[tuple[int, int], tuple[str, str]] = {}
    for node in element.findall(f"{W_NS}num"):
        reference = node.find(f"{W_NS}abstractNumId")
        levels = abstract.get(reference.get(f"{W_NS}val", "")) if reference is not None else None
        if not levels:
            continue
        num_id = int(node.get(f"{W_NS}numId", "0"))
        for ilvl, spec in levels.items():
            mapping[(num_id, ilvl)] = spec
    return mapping


def _numeral(value: int, fmt: str) -> str:
    alphabet = LETTER_FORMATS.get(fmt)
    if alphabet:
        return alphabet[(value - 1) % len(alphabet)]
    return str(value)


def _list_marker(spec: tuple[str, str], ilvl: int, counters: dict[int, int]) -> str:
    fmt, template = spec
    if fmt in BULLET_FORMATS:
        return "—"
    marker = template
    for level, count in counters.items():
        numeral = _numeral(count, fmt if level == ilvl else "decimal")
        marker = marker.replace(f"%{level + 1}", numeral)
    return re.sub(r"%\d", "", marker)


def _docx_num_ref(paragraph: Paragraph) -> tuple[int, int] | None:
    properties = paragraph._p.pPr
    numbering = properties.numPr if properties is not None else None
    if numbering is None or numbering.numId is None:
        return None
    ilvl = numbering.ilvl.val if numbering.ilvl is not None else 0
    return int(numbering.numId.val), int(ilvl or 0)


def _docx_list_level(paragraph: Paragraph) -> int | None:
    """Уровень пункта списка: сначала настоящая нумерация Word, потом отступ слева."""
    properties = paragraph._p.pPr
    numbering = properties.numPr if properties is not None else None
    if numbering is not None:
        raw_level = numbering.ilvl.val if numbering.ilvl is not None else 0
        return 1 + min(MAX_LIST_LEVEL - 1, int(raw_level or 0))
    indent = paragraph.paragraph_format.left_indent
    if indent is None:
        return None
    return 1 + min(MAX_LIST_LEVEL - 1, int(indent.pt // INDENT_STEP))


def _docx_element(
    paragraph: Paragraph, bbox: tuple[float, float, float, float], marker: str
) -> ParsedElement:
    style = paragraph.style.name.lower() if paragraph.style else ""
    text = f"{marker} {paragraph.text.strip()}".strip() if marker else paragraph.text.strip()
    list_level = _docx_list_level(paragraph)
    if marker or "list" in style or BULLET_RE.match(text):
        return ParsedElement("list", text, bbox, list_level or 1)
    if style.startswith("heading"):
        level_match = re.search(r"(\d+)$", style)
        return ParsedElement("heading", text, bbox, int(level_match.group(1)) if level_match else 1)
    if NUMBERED_RE.match(text) and len(text) <= HEADING_MAX_CHARS:
        return ParsedElement("heading", text, bbox, 1)
    if ORDERED_RE.match(text):
        return ParsedElement("list", text, bbox, list_level or 1)
    return ParsedElement("paragraph", text, bbox, None)


R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def _docx_images(
    document: Document, paragraph: Paragraph, index: int
) -> list[tuple[str, bytes]]:
    """Картинки абзаца: Word держит их отдельными частями, абзац ссылается по rId."""
    found: list[tuple[str, bytes]] = []
    for order, node in enumerate(paragraph._p.iter()):
        rel_id = node.get(R_EMBED)
        if not rel_id:
            continue
        part = document.part.related_parts.get(rel_id)
        if part is None:
            continue
        extension = Path(str(part.partname)).suffix.lstrip(".").lower() or "png"
        found.append((f"docx{index}-{order}.{extension}", part.blob))
    return found


def _docx_page(path: Path, owner: str = "") -> ParsedPage:
    document = Document(path)
    numbering = _docx_numbering(document)
    paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() or (owner and _docx_images(document, paragraph, 0))
    ]
    total = max(1, len(paragraphs))
    counters: dict[int, dict[int, int]] = {}

    elements: list[ParsedElement] = []
    for index, paragraph in enumerate(paragraphs):
        bbox = (0.0, index / total, 1.0, (index + 1) / total)
        if owner:
            for name, data in _docx_images(document, paragraph, index):
                elements.append(
                    ParsedElement(
                        "image",
                        IMAGE_PLACEHOLDER,
                        bbox,
                        None,
                        asset_path=store_material_asset(owner, name, data),
                    )
                )
        if not paragraph.text.strip():
            continue
        marker = ""
        reference = _docx_num_ref(paragraph)
        spec = numbering.get(reference) if reference is not None else None
        if reference is not None and spec is not None:
            num_id, ilvl = reference
            level_counts = counters.setdefault(num_id, {})
            level_counts[ilvl] = level_counts.get(ilvl, 0) + 1
            for deeper in [key for key in level_counts if key > ilvl]:
                del level_counts[deeper]
            marker = _list_marker(spec, ilvl, level_counts)
        elements.append(_docx_element(paragraph, bbox, marker))

    plain = "\n".join(element.text for element in elements if element.kind != "image")
    return ParsedPage(1, 1, 1, _markdown(elements), plain, "native", tuple(elements))


def _page_indices(
    document: fitz.Document, start_page: int, page_numbers: Sequence[int] | None
) -> Iterator[int]:
    """Какие страницы разбирать: продолжение с чекпоинта или явный список."""
    if page_numbers is None:
        yield from range(start_page - 1, len(document))
        return
    for page_number in page_numbers:
        if start_page <= page_number <= len(document):
            yield page_number - 1


def _text_layer_page(
    document: fitz.Document,
    page: fitz.Page,
    page_index: int,
    owner: str,
    mode: ParserMode,
    params: OcrRuntimeParams,
) -> ParsedPage:
    """Страница с готовым текстовым слоем: разметка плюс картинки со страницы."""
    legacy = _native_pdf_page(
        page, page_index + 1, owner, ocr_images=mode == ParserMode.FAST, params=params
    )
    try:
        parsed = parse_layout_page(document, page_index, owner)
    except (ValueError, KeyError, RuntimeError) as error:
        log.warning(
            "разметка страницы %s не удалась, откат на текстовый слой: %s",
            page_index + 1,
            error,
        )
        quality, confidence = _page_quality(legacy.elements, params.quality_threshold)
        return replace(
            legacy,
            quality=quality,
            confidence=confidence,
            diagnostics=(*legacy.diagnostics, "layout_fallback"),
        )
    # Схему разметчик теперь отдаёт сам (`picture`), а встроенный растр приходит
    # из текстового слоя. На одной и той же иллюстрации это два описания одного
    # объекта: без проверки перекрытия страница получала бы двойной фрагмент.
    layout_pictures = tuple(
        element.bbox for element in parsed.elements if element.kind == "image"
    )
    images = tuple(
        element
        for element in legacy.elements
        if element.kind == "image"
        and not any(_bbox_overlap(element.bbox, box) >= 0.6 for box in layout_pictures)
    )
    if not images:
        return parsed
    elements = _merge_native_and_images(parsed.elements, images)
    quality, confidence = _page_quality(elements, params.quality_threshold)
    return replace(
        parsed,
        markdown=_markdown(list(elements)),
        quality=quality,
        elements=elements,
        confidence=confidence,
    )


def _render_page(page: fitz.Page, params: OcrRuntimeParams) -> tuple[bytes, float]:
    """Растр страницы под распознавание и разрешение, с которым он снят."""
    scale, dpi = raster.render_scale(page, params.raster_scale)
    pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
    return pixmap.tobytes("png"), dpi


def _scanned_page(
    page: fitz.Page,
    page_index: int,
    owner: str,
    params: OcrRuntimeParams,
    recognizer: PageRecognizer | None,
) -> ParsedPage:
    """Страница без текстового слоя: целиком во внешнюю модель или в локальный OCR."""
    image, dpi = _render_page(page, params)
    if recognizer is not None:
        parsed = recognizer.recognize_page(
            image, page_index + 1, page.rect.width, page.rect.height
        )
        parsed = _attach_region_assets(page, parsed, owner)
        return replace(parsed, diagnostics=(*parsed.diagnostics, f"render_dpi:{dpi:.0f}"))
    with NamedTemporaryFile(suffix=".png", delete=False) as temporary:
        temporary.write(image)
        temporary_path = Path(temporary.name)
    try:
        parsed = paddle_fast.parse_image(
            temporary_path,
            page_index + 1,
            language=params.fast_language,
            ocr_version=params.fast_model_id,
            quality_threshold=params.quality_threshold,
            owner=owner,
        )
        return replace(parsed, diagnostics=(*parsed.diagnostics, f"render_dpi:{dpi:.0f}"))
    finally:
        temporary_path.unlink(missing_ok=True)


def _attach_region_assets(page: fitz.Page, parsed: ParsedPage, owner: str) -> ParsedPage:
    """Вырезать со страницы то, что модель прочитала, но показать не может.

    Прочитав страницу целиком, внешняя модель возвращает схему или график
    одной фразой: «блок-схема конечного автомата». Самой картинки в ответе нет
    и быть не может, поэтому в материале на её месте оставалась подпись без
    изображения. Координаты у модели при этом есть — по ним страница и режется
    локально, тем же вырезом, что уходит в модель в стратегии «Только то, что
    не читается». Формула и таблица режутся заодно: их LaTeX и Markdown
    просмотрщик показывает только тогда, когда они собираются, а оригинал
    нужен всегда.
    """
    if not owner:
        return parsed
    elements = list(parsed.elements)
    cropped = 0
    for index, element in enumerate(elements):
        if element.kind not in CROPPED_KINDS or element.asset_path or not element.bbox_reliable:
            continue
        name = f"p{parsed.page_number}-{element.kind}{index}.png"
        image = raster.region_image(page, element.bbox)
        elements[index] = replace(element, asset_path=store_material_asset(owner, name, image))
        cropped += 1
    if not cropped:
        return parsed
    return replace(
        parsed,
        elements=tuple(elements),
        diagnostics=(*parsed.diagnostics, f"cloud_crops:{cropped}"),
    )


def _recognized_regions(
    page: fitz.Page, parsed: ParsedPage, recognizer: PageRecognizer
) -> ParsedPage:
    """Дочитать внешней моделью только то, что текстовый слой не объясняет.

    Смысл режима: текст страницы уже есть и он точен — платить за его повторное
    распознавание незачем. Наружу уходят вырезы формул, схем и таблиц-картинок,
    а на их место встаёт то, что модель прочитала, с пометкой источника `vl`.
    """
    targets = [
        (index, element)
        for index, element in enumerate(parsed.elements)
        if element.kind in {"image", "formula"}
    ]
    if not targets:
        return parsed
    requests = [
        RegionRequest(index=index, kind=element.kind, image=raster.region_image(page, element.bbox))
        for index, element in targets
    ]
    answers = recognizer.recognize_regions(requests, parsed.page_number)
    recognized = {answer.index: answer for answer in answers}
    elements = list(parsed.elements)
    for index, element in targets:
        answer = recognized.get(index)
        if answer is None or not answer.text.strip():
            continue
        text = answer.text.strip()
        if answer.kind == "formula":
            # Инструкция вырезов просит LaTeX без обрамления (крупная страница
            # рядом уже даёт контекст, окружать $ там незачем) — обрамляем
            # здесь сами, иначе формула вернётся в текст сырым LaTeX и KaTeX
            # её не отрисует.
            text = wrap_bare_latex(text)
        elements[index] = replace(
            element,
            kind=answer.kind,
            text=text,
            confidence=answer.confidence,
            recognition_source="vl",
        )
    updated = tuple(elements)
    quality, confidence = _page_quality(updated)
    return replace(
        parsed,
        elements=updated,
        markdown=_markdown(list(updated)),
        plain_text="\n".join(item.text for item in updated if item.kind != "image"),
        quality=quality,
        confidence=confidence,
        diagnostics=(*parsed.diagnostics, f"cloud_regions:{len(requests)}"),
    )


def _pdf_pages(
    path: Path,
    mode: ParserMode,
    start_page: int,
    params: OcrRuntimeParams,
    page_numbers: Sequence[int] | None,
    owner: str,
    recognizer: PageRecognizer | None,
) -> Iterator[ParsedPage]:
    """Постраничный разбор PDF: у каждой страницы своя ветка по наличию текста."""
    document = fitz.open(path)
    whole_page = mode == ParserMode.CLOUD and params.cloud_strategy == "page"
    for page_index in _page_indices(document, start_page, page_numbers):
        page = document[page_index]
        if whole_page or not page.get_text("text").strip():
            yield _scanned_page(page, page_index, owner, params, recognizer)
            continue
        parsed = _text_layer_page(document, page, page_index, owner, mode, params)
        if recognizer is not None:
            parsed = _recognized_regions(page, parsed, recognizer)
        yield parsed


def _photo_page(
    path: Path,
    params: OcrRuntimeParams,
    owner: str,
    recognizer: PageRecognizer | None,
) -> ParsedPage:
    """Отдельная картинка (снимок страницы или скан) как единственная страница."""
    if recognizer is not None:
        with Image.open(path) as image:
            width, height = image.size
        return recognizer.recognize_page(path.read_bytes(), 1, float(width), float(height))
    return paddle_fast.parse_image(
        path,
        1,
        language=params.fast_language,
        ocr_version=params.fast_model_id,
        quality_threshold=params.quality_threshold,
        owner=owner,
    )


def iter_pages(
    path: Path,
    mode: ParserMode,
    start_page: int = 1,
    *,
    params: OcrRuntimeParams | None = None,
    page_numbers: Sequence[int] | None = None,
    recognizer: PageRecognizer | None = None,
) -> Iterator[ParsedPage]:
    """Разобрать файл постранично выбранным режимом.

    :param mode: режим распознавания; `CLOUD` работает только вместе с `recognizer`.
    :param start_page: с какой страницы продолжать после чекпоинта воркера.
    :param page_numbers: явный список страниц, когда переразбирается часть материала.
    :param recognizer: порт внешней модели; `None` — разбор целиком локальный.
    """
    params = params or OcrRuntimeParams()
    suffix = path.suffix.lower()
    # Имя файла материала — его sha256, поэтому оно же служит папкой для картинок.
    owner = path.stem
    if suffix == ".pdf":
        yield from _pdf_pages(path, mode, start_page, params, page_numbers, owner, recognizer)
        return
    if start_page > 1:
        return
    if suffix in {".jpg", ".jpeg", ".png"}:
        yield _photo_page(path, params, owner, recognizer)
    elif suffix == ".docx":
        yield _docx_page(path, owner)
    elif suffix in {".mp3", ".wav", ".m4a", ".ogg", ".flac"}:
        yield parse_audio(path)
    else:
        yield _text_page(path)
